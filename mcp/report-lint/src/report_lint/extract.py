"""文档抽取：把 docx / pdf / md 统一成一份可检索的文本与结构描述。

抽取只读，不修改源文件。缺少解析依赖时明确报错，不静默返回空结果。
"""
from __future__ import annotations

import pathlib
import re
from dataclasses import dataclass, field


class ExtractError(RuntimeError):
    """抽取失败，带错误码，便于 CLI 与 MCP 层区分处理。"""

    def __init__(self, code: str, message: str, hint: str = "") -> None:
        super().__init__(message)
        self.code = code
        self.hint = hint


@dataclass
class Document:
    path: str
    kind: str
    headings: list[dict] = field(default_factory=list)
    paragraphs: list[str] = field(default_factory=list)
    tables: list[list[list[str]]] = field(default_factory=list)
    pages: int | None = None
    text: str = ""

    def as_dict(self) -> dict:
        return {
            "path": self.path,
            "kind": self.kind,
            "pages": self.pages,
            "headings": self.headings,
            "paragraphs": len(self.paragraphs),
            "tables": len(self.tables),
            "chars": len(self.text),
        }


_HEADING_RE = re.compile(r"^Heading\s+(\d)$", re.IGNORECASE)


def load(path: str | pathlib.Path) -> Document:
    p = pathlib.Path(path)
    if not p.exists():
        raise ExtractError("E_NOT_FOUND", f"文件不存在：{p}", "确认路径与文件名")
    suffix = p.suffix.lower()
    if suffix == ".docx":
        return _load_docx(p)
    if suffix == ".pdf":
        return _load_pdf(p)
    if suffix in {".md", ".markdown", ".txt"}:
        return _load_text(p)
    raise ExtractError(
        "E_UNSUPPORTED",
        f"暂不支持 {suffix} 格式",
        "支持 .docx / .pdf / .md / .txt；旧版 .doc 请先另存为 .docx",
    )


def _iter_block_items(doc):
    """按文档顺序产出段落与表格。

    python-docx 的 doc.paragraphs / doc.tables 是分开取的，直接拼接会让
    「VULN 属性表」全部落到全文末尾，按条切分时就找不到风险等级字段。
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _load_docx(p: pathlib.Path) -> Document:
    try:
        from docx import Document as Docx
    except ImportError as exc:  # pragma: no cover - 环境缺依赖时给出可操作提示
        raise ExtractError(
            "E_DEPS_MISSING", "缺少 python-docx", "pip install python-docx"
        ) from exc

    doc = Docx(str(p))
    headings: list[dict] = []
    paragraphs: list[str] = []
    tables: list[list[list[str]]] = []
    blocks: list[str] = []

    for block in _iter_block_items(doc):
        if block.__class__.__name__ == "Paragraph":
            title = block.text.strip()
            if not title:
                continue
            paragraphs.append(title)
            blocks.append(title)
            try:
                style = block.style.name or ""
            except Exception:
                style = ""
            match = _HEADING_RE.match(style)
            if match:
                headings.append({"level": int(match.group(1)), "title": title})
            continue

        rows = [[cell.text.strip() for cell in row.cells] for row in block.rows]
        tables.append(rows)
        for row in rows:
            blocks.append(" | ".join(row))

    return Document(
        path=str(p),
        kind="docx",
        headings=headings,
        paragraphs=paragraphs,
        tables=tables,
        pages=None,
        text="\n".join(blocks),
    )


def _load_pdf(p: pathlib.Path) -> Document:
    try:
        import pymupdf
    except ImportError as exc:  # pragma: no cover
        raise ExtractError(
            "E_DEPS_MISSING", "缺少 pymupdf", "pip install pymupdf"
        ) from exc

    try:
        pdf = pymupdf.open(str(p))
    except Exception as exc:
        raise ExtractError(
            "E_ENCRYPTED", f"PDF 无法打开：{exc}", "若已加密请先解密另存"
        ) from exc

    pages = [page.get_text() for page in pdf]
    text = "\n".join(pages).strip()
    if not text:
        raise ExtractError(
            "E_NO_TEXT_LAYER",
            "PDF 没有文本层，可能是扫描件",
            "不做 OCR 猜测；请提供 docx 或带文本层的 PDF",
        )

    headings = []
    for line in text.splitlines():
        bare = line.strip()
        # 只做粗识别：章节标题通常是「1 概述」「3.2.1 VULN-01 ...」这类短行
        if re.match(r"^(\d+(\.\d+)*)\s+\S", bare) and len(bare) <= 60:
            level = bare.split()[0].count(".") + 2
            headings.append({"level": level, "title": bare})

    return Document(
        path=str(p),
        kind="pdf",
        headings=headings,
        paragraphs=[ln.strip() for ln in text.splitlines() if ln.strip()],
        tables=[],
        pages=len(pages),
        text=text,
    )


def _load_text(p: pathlib.Path) -> Document:
    text = p.read_text(encoding="utf-8", errors="replace")
    headings = []
    for line in text.splitlines():
        match = re.match(r"^(#{1,4})\s+(.+)$", line.strip())
        if match:
            headings.append({"level": len(match.group(1)), "title": match.group(2).strip()})
    return Document(
        path=str(p),
        kind="md" if p.suffix.lower() != ".txt" else "txt",
        headings=headings,
        paragraphs=[ln.strip() for ln in text.splitlines() if ln.strip()],
        tables=[],
        pages=None,
        text=text,
    )
